import io
import json
from unittest.mock import patch

import pytest
from docx import Document
from docx.oxml import OxmlElement

from app.config import Settings
from app.detectors.context_rules import is_contract_label, is_technical_size, signatures
from app.detectors.ollama_detector import OllamaDetector, OllamaError, Review
from app.files.office import _docx_targets, _replace_in_paragraph
from app.pipeline import AnonymizerPipeline
from app.schemas import Finding


def test_docx_tabs_and_drawing_do_not_leave_original_name():
    document = Document()
    paragraph = document.add_paragraph()
    run = paragraph.add_run('Главный инженер\t\tИ. В. Запорожец')
    drawing = OxmlElement('w:drawing')
    run._r.append(drawing)
    paragraph.add_run(' / подпись').bold = True
    expected = 'Главный инженер\t\t[ФИО] / подпись'
    _replace_in_paragraph(paragraph, expected)
    assert paragraph.text == expected
    assert drawing.getparent() is run._r
    buffer = io.BytesIO()
    document.save(buffer)
    reopened = Document(io.BytesIO(buffer.getvalue()))
    assert reopened.paragraphs[0].text == expected
    assert 'Запорожец' not in reopened.element.xml


def test_multiple_replacements_across_runs():
    document = Document()
    paragraph = document.add_paragraph('Иван ')
    paragraph.add_run('Петров, телефон: ')
    paragraph.add_run('+7 912 345-67-89. Иван Петров.')
    expected = '[ФИО], телефон: [ТЕЛЕФОН]. [ФИО].'
    _replace_in_paragraph(paragraph, expected)
    assert paragraph.text == expected


def test_docx_traversal_keeps_all_xml_elements_alive():
    document = Document()
    for index in range(200):
        document.add_table(rows=1, cols=1).cell(0, 0).text = f'Ячейка {index}'
    first = [p.text for p, _ in _docx_targets(document)]
    assert first == [f'Ячейка {index}' for index in range(200)]
    assert [p.text for p, _ in _docx_targets(document)] == first


@pytest.mark.parametrize('text', [
    'Главный инженерИ. В. Запорожец',
    'Начальник юридического отдела\tИ. В. Жданова',
    'Подписант: Ершова А. В.',
    'Руководитель контрактной службы\tЕ. А. Морозова',
    'Специалист по закупкам\tГ. С. Маркова',
    '________________ В. В. Агафонов',
])
def test_signature_rules(text):
    found = signatures(text)
    assert len(found) == 1
    assert found[0].entity == 'PERSON'
    assert 'инженер' not in found[0].extract(text)


def test_du_is_excluded_only_in_dimension_context():
    for text in ['Кран шаровый Ду=16', 'Инженер Ду подписал акт']:
        start = text.index('Ду')
        finding = Finding(entity='PERSON', start=start, end=start + 2)
        assert is_technical_size(text, finding) == ('=16' in text)


def test_contract_roles_are_not_persons_but_surnames_are():
    for text in ['Покупателю', 'Поставщика', 'Заказчиком', 'Документации', 'Петров']:
        finding = Finding(entity='PERSON', start=0, end=len(text))
        assert is_contract_label(text, finding) == (text != 'Петров')


def test_equipment_with_dimension_is_not_a_person():
    text = 'Резьба стальная Ду 15 мм'
    assert is_technical_size(text, Finding(entity='PERSON', start=0, end=len(text)))


def test_review_removes_false_person_adds_missing_name_keeps_regex():
    text = 'Ду=16. Инженер И. В. Запорожец. ИНН 7707083893'
    findings = [Finding(entity='PERSON', start=0, end=2, source='presidio'),
                Finding(entity='INN_RU', start=text.index('770'), end=len(text), source='regex')]
    detector = OllamaDetector(Settings())
    response = Review.model_validate({'decisions': [{'id': 0, 'keep': False}],
                                     'entities': [{'text': 'И. В. Запорожец', 'entity': 'PERSON'}]})
    with patch.object(detector, '_review', return_value=response):
        result = detector.refine(text, findings, None)
    assert [(f.entity, f.extract(text)) for f in result] == [
        ('INN_RU', '7707083893'), ('PERSON', 'И. В. Запорожец')]


@pytest.mark.parametrize('url', ['https://example.com', 'http://127.0.0.1.example.com', 'http://user@localhost:11434'])
def test_ollama_rejects_remote_addresses(url):
    with pytest.raises(ValueError):
        OllamaDetector(Settings(ollama_url=url))


def test_container_host_requires_explicit_enablement():
    with pytest.raises(ValueError):
        OllamaDetector(Settings(ollama_url='http://ollama:11434'))
    detector = OllamaDetector(Settings(ollama_url='http://ollama:11434',
                                      ollama_allow_container_host=True))
    assert detector.settings.ollama_url == 'http://ollama:11434'


@pytest.mark.parametrize('url', [
    'http://ollama.example.com:11434', 'http://ollama:8080', 'http://ollama:11434/proxy',
])
def test_container_mode_does_not_allow_arbitrary_servers(url):
    with pytest.raises(ValueError):
        OllamaDetector(Settings(ollama_url=url, ollama_allow_container_host=True))


@pytest.mark.parametrize('content', [
    '{"decisions": [], "entities": []}',
    '{"decisions": [{"id": 0, "keep": true}], "entities": [{"text": "Выдуманное имя", "entity": "PERSON"}]}',
])
def test_invalid_review_fails_instead_of_returning_file(content):
    detector = OllamaDetector(Settings())
    with patch.object(detector._http, 'open') as opened:
        opened.return_value.__enter__.return_value = io.BytesIO(
            ('{"message":{"content":' + json.dumps(content) + '}}').encode())
        with pytest.raises(OllamaError):
            detector._review('Покупатель', [{'id': 0}], {'PERSON'})


def test_regex_only_request_does_not_call_ollama():
    detector = OllamaDetector(Settings())
    with patch.object(detector, '_review', side_effect=AssertionError('Unexpected LLM call')):
        assert detector.refine('ИНН 7707083893', [], {'INN_RU'}) == []


def test_llm_cannot_replace_a_regex_number_with_a_person():
    text = 'ИНН 7707083893'
    number = Finding(entity='INN_RU', start=4, end=len(text), source='regex', score=0.7)
    detector = OllamaDetector(Settings())
    response = Review.model_validate({'decisions': [], 'entities': [
        {'text': text, 'entity': 'PERSON'}]})
    with patch.object(detector, '_review', return_value=response):
        assert detector.refine(text, [number], None) == [number]


def test_review_retries_invalid_substring_and_accepts_verified_response():
    detector = OllamaDetector(Settings())
    replies = [
        {'decisions': [], 'entities': [{'text': 'Выдуманное имя', 'entity': 'PERSON'}]},
        {'decisions': [], 'entities': [{'text': 'Иван Петров', 'entity': 'PERSON'}]},
    ]
    def respond(*args, **kwargs):
        return io.BytesIO(json.dumps({'message': {'content': json.dumps(replies.pop(0))}}).encode())
    with patch.object(detector._http, 'open', side_effect=respond):
        review = detector._review('Иван Петров', [], {'PERSON'})
    assert review.entities[0].text == 'Иван Петров'
    assert not replies


def test_context_rules_in_real_pipeline():
    pipeline = AnonymizerPipeline(Settings(ollama_enabled=False))
    text = 'Главный инженер\tИ. В. Запорожец\nКран шаровый Ду=16, PN16.'
    result = pipeline.anonymize(text, entities=['PERSON'])
    assert 'Запорожец' not in result.text
    assert 'Ду=16' in result.text


def test_api_returns_503_when_local_review_fails(monkeypatch):
    from fastapi.testclient import TestClient

    from app import main

    monkeypatch.setattr(main.settings, 'api_key', '')
    monkeypatch.setattr(main.settings, 'env', 'development')
    with (
        TestClient(main.app) as client,
        patch.object(main.pipeline, 'anonymize', side_effect=OllamaError('Проверка не завершена')),
    ):
        response = client.post('/api/v1/anonymize/text', json={'text': 'Иван Петров'})
    assert response.status_code == 503
    assert response.json() == {'detail': 'Проверка не завершена'}
    assert 'findings' not in response.json()
