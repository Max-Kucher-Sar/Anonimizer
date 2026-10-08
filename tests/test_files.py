"""Интеграционные тесты файлового анонимайзера на настоящих детекторах."""

from __future__ import annotations

import io
import json

import pytest
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from app.config import get_settings
from app.files.pipeline import FilePipeline
from app.pipeline import AnonymizerPipeline
from tests.file_fixtures import (
    EMAIL,
    INN,
    PASSPORT,
    PERSON,
    PHONE,
    make_broken_json,
    make_csv,
    make_docx,
    make_docx_split_name,
    make_json,
    make_pptx,
    make_txt,
    make_txt_cp1251,
    make_xlsx,
    make_xlsx_numbers,
)


@pytest.fixture(scope="module")
def pipeline() -> AnonymizerPipeline:
    return AnonymizerPipeline(get_settings())


@pytest.fixture(scope="module")
def files(pipeline) -> FilePipeline:
    return FilePipeline(pipeline, get_settings())


def run(files: FilePipeline, payload: bytes, name: str, **kwargs):
    return files.anonymize(payload=payload, filename=name, **kwargs)


def as_text(payload: bytes) -> str:
    return payload.decode("utf-8")


class TestPlainText:
    def test_finds_and_replaces(self, files: FilePipeline) -> None:
        result = run(files, make_txt(), "договор.txt")

        text = as_text(result.payload)
        assert PASSPORT not in text
        assert INN not in text
        assert EMAIL not in text
        assert PHONE not in text
        assert "Без персональных данных" in text

    def test_keeps_structure(self, files: FilePipeline) -> None:
        """Структура документа обязана выжить: иначе файл нечитаем."""

        result = run(files, make_txt(), "договор.txt")

        assert as_text(result.payload).count("\n") == as_text(make_txt()).count("\n")

    def test_trailing_newline_preserved(self, files: FilePipeline) -> None:
        payload = "Паспорт 45 08 123456\n"

        result = run(files, payload.encode(), "a.txt")

        assert as_text(result.payload).endswith("\n")

    def test_units(self, files: FilePipeline) -> None:
        assert run(files, make_txt(), "a.txt").units == 1

    def test_media_type(self, files: FilePipeline) -> None:
        assert run(files, make_txt(), "a.txt").media_type == "text/plain"

    def test_counts_sum(self, files: FilePipeline) -> None:
        result = run(files, make_txt(), "a.txt")

        assert sum(result.counts().values()) == len(result.findings)

    def test_cp1251_is_read_and_warned(self, files: FilePipeline) -> None:
        """Файл в windows-1251 читается, но о не-UTF8 честно предупреждаем."""

        result = run(files, make_txt_cp1251(), "a.txt")

        assert any("не в UTF-8" in warning for warning in result.warnings)

    def test_entities_filter(self, files: FilePipeline) -> None:
        result = run(files, make_txt(), "a.txt", entities=["EMAIL_ADDRESS"])

        assert [finding.entity for finding in result.findings] == ["EMAIL_ADDRESS"]
        assert PASSPORT in as_text(result.payload)

    def test_empty_file_rejected(self, files: FilePipeline) -> None:
        with pytest.raises(ValueError, match="пустой"):
            run(files, b"", "a.txt")


class TestCsv:
    def test_redacts_values(self, files: FilePipeline) -> None:
        result = run(files, make_csv(), "clients.csv")

        text = as_text(result.payload)
        assert PERSON not in text
        assert PHONE not in text
        assert "подписан" in text

    def test_column_count_stable(self, files: FilePipeline) -> None:
        """Заглушка не должна разъехаться по колонкам."""

        result = run(files, make_csv(), "clients.csv")

        rows = [line for line in as_text(result.payload).split("\n") if line]
        assert all(len(row.split(",")) == 4 for row in rows)

    def test_media_type(self, files: FilePipeline) -> None:
        assert run(files, make_csv(), "a.csv").media_type == "text/csv"


class TestJson:
    def test_only_strings_touched(self, files: FilePipeline) -> None:
        result = run(files, make_json(), "a.json")

        document = json.loads(as_text(result.payload))
        assert document["client"]["name"] == "[ФИО]"
        assert document["contacts"][0]["email"] == "[EMAIL]"

    def test_passport_value_is_replaced(self, files: FilePipeline) -> None:
        """Важно, что значение заменено, а не как именно оно названо.

        Без подписи «Паспорт» регулярка паспорта не срабатывает и строку
        забирает распознаватель телефонов — это поведение исходных детекторов.
        Тип в отчёте может оказаться не тем, а вот замена обязана быть.
        """
        result = run(files, make_json(), "a.json")

        document = json.loads(as_text(result.payload))
        assert PASSPORT not in document["passport"]
        assert document["passport"].startswith("[")

    def test_numeric_pii_is_zeroed(self, files: FilePipeline) -> None:
        """ИНН, записанный числом, тоже ПДн — раньше он оставался как есть.

        Обнуляется, а не подменяется строкой: тип обязан остаться прежним,
        иначе документ перестанет соответствовать своей схеме.
        """
        result = run(files, make_json(), "a.json")

        document = json.loads(as_text(result.payload))
        assert document["client"]["inn"] == 0
        assert isinstance(document["client"]["inn"], int)

    def test_bool_and_null_untouched(self, files: FilePipeline) -> None:
        result = run(files, make_json(), "a.json")

        document = json.loads(as_text(result.payload))
        assert document["client"]["active"] is True

    def test_numbers_zeroed_warning(self, files: FilePipeline) -> None:
        result = run(files, make_json(), "a.json")

        assert any("обнулены" in warning for warning in result.warnings)

    def test_json_without_numbers_has_no_warning(self, files: FilePipeline) -> None:
        payload = json.dumps({"name": "Иванов Иван Иванович"}, ensure_ascii=False).encode()

        result = run(files, payload, "a.json")

        assert not any("обнулены" in warning for warning in result.warnings)

    def test_output_is_valid_json(self, files: FilePipeline) -> None:
        result = run(files, make_json(), "a.json")

        json.loads(as_text(result.payload))

    def test_structure_preserved(self, files: FilePipeline) -> None:
        result = run(files, make_json(), "a.json")

        assert json.loads(as_text(result.payload)).keys() == json.loads(
            as_text(make_json())
        ).keys()

    def test_reports_json_paths(self, files: FilePipeline) -> None:
        result = run(files, make_json(), "a.json")

        assert any(finding.location.startswith("client.") for finding in result.findings)

    def test_broken_json_rejected(self, files: FilePipeline) -> None:
        with pytest.raises(ValueError, match="не является корректным JSON"):
            run(files, make_broken_json(), "a.json")

    def test_media_type(self, files: FilePipeline) -> None:
        assert run(files, make_json(), "a.json").media_type == "application/json"


class TestDocx:
    @staticmethod
    def text_of(payload: bytes) -> str:
        document = Document(io.BytesIO(payload))
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.extend(paragraph.text for paragraph in cell.paragraphs)
        return "\n".join(parts)

    def test_redacts_paragraphs_and_tables(self, files: FilePipeline) -> None:
        result = run(files, make_docx(), "a.docx")

        text = self.text_of(result.payload)
        assert PASSPORT not in text
        assert INN not in text
        assert PERSON not in text

    def test_keeps_heading(self, files: FilePipeline) -> None:
        result = run(files, make_docx(), "a.docx")

        assert "Договор поставки" in self.text_of(result.payload)

    def test_split_name_across_paragraphs(self, files: FilePipeline) -> None:
        """Три абзаца с частями ФИО: склейка должна собрать их в сущность."""
        result = run(files, make_docx_split_name(), "a.docx")

        text = self.text_of(result.payload)
        assert "Иванов" not in text
        assert "Иванович" not in text

    def test_reports_paragraph_numbers(self, files: FilePipeline) -> None:
        result = run(files, make_docx(), "a.docx")

        assert any(finding.location.startswith("абзац") for finding in result.findings)
        assert any(finding.location.startswith("таблица") for finding in result.findings)

    def test_units(self, files: FilePipeline) -> None:
        assert run(files, make_docx(), "a.docx").units >= 4

    def test_media_type(self, files: FilePipeline) -> None:
        assert "wordprocessingml" in run(files, make_docx(), "a.docx").media_type

    def test_broken_docx_rejected(self, files: FilePipeline) -> None:
        with pytest.raises(ValueError, match="docx"):
            run(files, b"not a docx", "a.docx")


class TestXlsx:
    @staticmethod
    def cell_values(payload: bytes) -> dict[str, object]:
        workbook = load_workbook(io.BytesIO(payload), data_only=False)
        values: dict[str, object] = {}
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    values[f"{sheet.title}!{cell.coordinate}"] = cell.value
        return values

    def test_redacts_cells(self, files: FilePipeline) -> None:
        result = run(files, make_xlsx(), "a.xlsx")

        values = self.cell_values(result.payload)
        assert PERSON not in str(values["Клиенты!A2"])
        assert PHONE not in str(values["Клиенты!B2"])
        assert INN not in str(values["Клиенты!A3"])
        assert EMAIL not in str(values["Архив!A1"])

    def test_formula_kept(self, files: FilePipeline) -> None:
        """Формулы не трогаем: это не текст, а вычисление."""

        result = run(files, make_xlsx(), "a.xlsx")

        assert self.cell_values(result.payload)["Клиенты!B3"] == "=B2"

    def test_formula_warning(self, files: FilePipeline) -> None:
        result = run(files, make_xlsx(), "a.xlsx")

        assert any("Формулы" in warning for warning in result.warnings)

    def test_headers_kept(self, files: FilePipeline) -> None:
        result = run(files, make_xlsx(), "a.xlsx")

        assert self.cell_values(result.payload)["Клиенты!A1"] == "ФИО"

    def test_sheet_count_preserved(self, files: FilePipeline) -> None:
        result = run(files, make_xlsx(), "a.xlsx")

        workbook = load_workbook(io.BytesIO(result.payload))
        assert workbook.sheetnames == ["Клиенты", "Архив"]

    def test_reports_cell_addresses(self, files: FilePipeline) -> None:
        result = run(files, make_xlsx(), "a.xlsx")

        assert any("!" in finding.location for finding in result.findings)

    def test_units(self, files: FilePipeline) -> None:
        """Считаются все непустые ячейки, включая заголовки и числа."""
        assert run(files, make_xlsx(), "a.xlsx").units == 6

    def test_numeric_pii_is_zeroed(self, files: FilePipeline) -> None:
        result = run(files, make_xlsx_numbers(), "a.xlsx")
        sheet = load_workbook(io.BytesIO(result.payload))["Клиенты"]

        assert sheet["A2"].value == 0

    def test_number_without_pii_keeps_type(self, files: FilePipeline) -> None:
        """Обычное число не должно превратиться в строку из-за подстановки."""
        result = run(files, make_xlsx_numbers(), "a.xlsx")
        sheet = load_workbook(io.BytesIO(result.payload))["Клиенты"]

        assert sheet["B2"].value == 1500.5
        assert isinstance(sheet["B2"].value, float)

    def test_number_warning(self, files: FilePipeline) -> None:
        result = run(files, make_xlsx_numbers(), "a.xlsx")

        assert any("обнулены" in warning for warning in result.warnings)

    def test_media_type(self, files: FilePipeline) -> None:
        assert "spreadsheetml" in run(files, make_xlsx(), "a.xlsx").media_type


class TestPptx:
    @staticmethod
    def text_of(payload: bytes) -> str:
        presentation = Presentation(io.BytesIO(payload))
        parts = []
        for slide in presentation.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    parts.append(shape.text_frame.text)
        return "\n".join(parts)

    def test_redacts_slides(self, files: FilePipeline) -> None:
        result = run(files, make_pptx(), "a.pptx")

        text = self.text_of(result.payload)
        assert PERSON not in text
        assert PHONE not in text

    def test_keeps_title(self, files: FilePipeline) -> None:
        result = run(files, make_pptx(), "a.pptx")

        assert "Отчёт по клиентам" in self.text_of(result.payload)

    def test_reports_slide_and_shape(self, files: FilePipeline) -> None:
        result = run(files, make_pptx(), "a.pptx")

        assert any(finding.location.startswith("слайд") for finding in result.findings)

    def test_media_type(self, files: FilePipeline) -> None:
        assert "presentationml" in run(files, make_pptx(), "a.pptx").media_type

    def test_broken_pptx_rejected(self, files: FilePipeline) -> None:
        with pytest.raises(ValueError, match="pptx"):
            run(files, b"not a pptx", "a.pptx")


class TestDispatch:
    @pytest.mark.parametrize(
        ("name", "payload"),
        [
            ("a.exe", b"x"),
            ("a", b"x"),
            ("a.doc", b"x"),
            ("a.pdf.exe", b"x"),
        ],
    )
    def test_rejects_unknown_format(self, files: FilePipeline, name: str, payload: bytes) -> None:
        with pytest.raises(ValueError, match="не поддерживается"):
            run(files, payload, name)

    def test_extension_is_case_insensitive(self, files: FilePipeline) -> None:
        result = run(files, make_txt(), "ДОГОВОР.TXT")

        assert PASSPORT not in as_text(result.payload)

    def test_size_limit(self, pipeline: AnonymizerPipeline) -> None:
        # Копия, а не мутация: get_settings закеширован, и правка разъехалась бы
        # на все остальные тесты.
        tiny = FilePipeline(
            pipeline, get_settings().model_copy(update={"max_file_bytes": 64})
        )

        with pytest.raises(ValueError, match="больше"):
            tiny.anonymize(payload=make_txt(), filename="a.txt")

    def test_fast_mode_works(self, files: FilePipeline) -> None:
        result = run(files, make_txt(), "a.txt")

        assert result.findings
